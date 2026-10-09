"""Parser, upload pin, and dashboard read for Outbound Productivity.

Synthetic workbooks only. The anonymous golden ratios are the expected totals.
"""
from __future__ import annotations

import io
import uuid
from datetime import datetime, time
from types import SimpleNamespace

import pandas as pd
import pytest
from openpyxl import Workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from api.dependencies import clear_serialization_cache, serialize_performance_record
from config.loader import load_team_config
from data_cleaning.standard_mappings import calculate_grade
from exports.report_exporter import ReportExporter
from models.models import (
    Employee,
    EmployeeUploadBatch,
    EvaluationRevision,
    EvaluationScope,
    KPIValue,
    PerformanceRecord,
    Team,
    TeamConfigurationVersion,
    TeamKPIConfig,
    UploadLog,
)
from models.schemas import EvaluationData, PerformanceRecord as SchemaRecord
from processors.excel_processor import ExcelProcessor
from services.dashboard_record_service import DashboardRecordService
from services.evaluation.access import TargetConflict
from services.kpi_aggregation import aggregate_kpi_metric
from services.outbound_period_basis import (
    OutboundUnsupportedApprovedBinding,
    assess_monthly_actuals,
    period_capability,
)
from services.seeding_service import DatabaseSeeder, UploadProcessingError
from Data_Cleaning_Teams.outbound import process_outbound

AUGUST = {
    "Booking": 0.3008083140877598,
    "Attendance": 0.46065259117082535,
    "Other": 0.568241469816273,
    "Quality": 0.98,
    "Productivity": 0.7780694444444445,
}
JULY = {
    "Booking": 0.1703551209469892,
    "Attendance": 0.552870090634441,
    "Other": 0.5724808485562758,
    "Quality": 0.96,
}
TARGETS = {"Booking": 0.3, "Attendance": 0.65, "Other": 0.75, "Quality": 0.95, "Productivity": 0.8}
AUGUST_RATIO = 0.7982419863297692
JULY_RATIO = 0.8285143792679492


@pytest.fixture(autouse=True)
def _no_external_cache(monkeypatch):
    monkeypatch.setattr("services.cache_service.redis_client", None, raising=False)
    monkeypatch.setattr("services.cache_invalidation_service.redis_client", None, raising=False)
    monkeypatch.setattr(
        "services.cache_invalidation_service.CacheInvalidationService.bump_data_version",
        staticmethod(lambda: 0),
    )
    monkeypatch.setattr(
        "services.cache_invalidation_service.CacheInvalidationService.bump_config_version",
        staticmethod(lambda: 0),
    )


@pytest.fixture()
def db():
    engine = create_engine(
        "sqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    from config.database import Base

    Base.metadata.create_all(
        bind=engine,
        tables=[
            Team.__table__,
            Employee.__table__,
            PerformanceRecord.__table__,
            KPIValue.__table__,
            TeamKPIConfig.__table__,
            TeamConfigurationVersion.__table__,
            EvaluationScope.__table__,
            EvaluationRevision.__table__,
            EmployeeUploadBatch.__table__,
            UploadLog.__table__,
        ],
    )
    session = sessionmaker(bind=engine, autoflush=False, autocommit=False)()
    try:
        yield session
    finally:
        session.close()


def _workbook(when: datetime, actuals: dict, *, productivity: float | None, role: str | None, sheet: str) -> bytes:
    headers = [
        "SGHCode", "EnglishName", "Date", "Status", "Region", "AHT", "Available Time",
        "A.Booking%", "T.Booking%", "A.Attend%", "T.Attend%",
        "A.Reachability%", "T.Reachability%", "A.QualityScore", "T.Quality%",
        "Performance Grade", "Productivity", "T.Productivity%", "Scratch Note",
    ]
    if role is not None:
        headers.insert(2, "Role")
    values = [
        "ANON-1", "Anonymous Agent", when, "Active", "EGY", time(0, 2, 30), 0.25,
        actuals["Booking"], TARGETS["Booking"], actuals["Attendance"], TARGETS["Attendance"],
        actuals["Other"], TARGETS["Other"], actuals["Quality"], TARGETS["Quality"],
        "C",
        "" if productivity is None else productivity,
        "" if productivity is None else TARGETS["Productivity"],
        "drop-me",
    ]
    if role is not None:
        values.insert(2, role)
    book = Workbook()
    sheet_obj = book.active
    sheet_obj.title = sheet
    sheet_obj.append(headers)
    sheet_obj.append(values)
    attend_col = headers.index("T.Attend%") + 1
    attend = sheet_obj.cell(2, attend_col)
    attend.value = TARGETS["Attendance"]
    attend.number_format = "m/d/yy h:mm"
    if productivity is None:
        for name in ("Productivity", "T.Productivity%"):
            column = headers.index(name) + 1
            sheet_obj.cell(2, column).value = None
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _august_bytes(**overrides) -> bytes:
    options = {
        "when": datetime(2026, 8, 15),
        "actuals": AUGUST,
        "productivity": AUGUST["Productivity"],
        "role": "Emp",
        "sheet": "Outbound Offshore Call Center",
    }
    options.update(overrides)
    return _workbook(**options)


def _seeder():
    seeder = DatabaseSeeder()
    seeder.performance_repo.get_all = lambda: []
    seeder.performance_repo.save_all = lambda *_args, **_kwargs: None
    return seeder


def _bind(db, monkeypatch):
    monkeypatch.setattr(
        "services.seeding_service.SessionLocal",
        sessionmaker(bind=db.get_bind(), autoflush=False, autocommit=False),
    )


def _by_key(rows):
    return {row.kpi_key: row for row in rows}


def _delete_record(db, stored) -> None:
    """Remove a record without the ORM nulling the KPI foreign key."""
    record_id = stored.id
    for row in list(stored.kpi_values):
        db.expunge(row)
    db.expunge(stored)
    db.query(KPIValue).filter(KPIValue.record_id == record_id).delete(synchronize_session=False)
    db.query(PerformanceRecord).filter(PerformanceRecord.id == record_id).delete(synchronize_session=False)
    db.commit()


def test_period_capability_keeps_july_and_august_separate_from_later_months():
    july = period_capability(2026, 7)
    august = period_capability(2026, "August")
    september = period_capability(2026, 9)
    other_august = period_capability(2027, 8)

    assert [line["weight"] for line in july["lines"]] == [0.70, 0.10, 0.10, 0.10]
    assert july["scored_keys"] == ["Attendance", "Booking", "Quality", "Other"]
    assert august["lines"][1]["kpi_key"] == "Attendance"
    assert august["lines"][1]["weight"] == 0.60
    assert august["scored_keys"][-1] == "Productivity"
    assert august["lines"][-1]["target"] == 0.8
    assert july["catalog_template_supported"] is True
    assert august["catalog_template_supported"] is True
    assert september["approved_binding"] == "refuse"
    assert september["catalog_template_supported"] is False
    assert september["infer_from_august"] is False
    assert other_august["status"] == "unconfigured"
    assert september["lines"] == other_august["lines"]
    assert "Productivity" not in september["scored_keys"]
    missing = assess_monthly_actuals(2026, 8, {key: AUGUST[key] for key in ("Booking", "Attendance", "Other", "Quality")})
    assert missing["sufficient"] is False
    assert missing["accepted_as_five_kpi_monthly_evidence"] is False
    assert missing["historical_score_readable"] is True
    assert missing["disposition"] == "fail_visible_missing_source_actual"
    july_evidence = assess_monthly_actuals(2026, 7, {"Booking": JULY["Booking"], "Attendance": JULY["Attendance"], "Other": JULY["Other"], "Quality": JULY["Quality"]})
    assert july_evidence["sufficient"] is True
    assert july_evidence["accepted_as_five_kpi_monthly_evidence"] is False


def test_format_22_percent_serial_survives_the_real_parser_and_crop():
    blob = _august_bytes()
    frame = ExcelProcessor().process_sheet_outbound(pd.ExcelFile(io.BytesIO(blob)))
    attend = frame.loc[0, "T.Attend%"]
    assert not isinstance(attend, time)
    assert float(attend) == 0.65
    assert frame.loc[0, "Date"].month == 8 and frame.loc[0, "Date"].year == 2026
    assert float(frame.loc[0, "AHT_Minutes"]) == pytest.approx(2.5)
    assert float(frame.loc[0, "Productivity"]) == pytest.approx(AUGUST["Productivity"])
    assert "ScratchNote" not in frame.columns
    assert float(frame.loc[0, "AvailableTime"]) == 0.25
    assert float(frame.loc[0, "Productivity"]) != 0.25

    bare = _august_bytes(role=None)
    legacy = process_outbound(bare)
    legacy_attend = legacy.loc[0, "T.Attend%"]
    assert not isinstance(legacy_attend, time)
    assert float(legacy_attend) == 0.65
    assert "Productivity" in legacy.columns


def test_august_upload_dry_run_matches_commit_and_keeps_the_five_kpi_total(db, monkeypatch):
    _bind(db, monkeypatch)
    seeder = _seeder()
    blob = _august_bytes()
    preview = seeder.process_uploaded_file("august.xlsx", blob, dry_run=True)
    db.rollback()
    assert db.query(PerformanceRecord).count() == 0
    committed = seeder.process_uploaded_file("august.xlsx", blob, dry_run=False)
    assert preview["scored_rows"] == committed["scored_rows"]
    assert preview["records_imported"] == committed["records_imported"] == 1

    stored = db.query(PerformanceRecord).one()
    kpis = _by_key(stored.kpi_values)
    assert set(kpis) == {"Booking", "Attendance", "Quality", "Other", "Productivity"}
    assert float(kpis["Attendance"].weight_applied) == pytest.approx(0.60)
    assert float(kpis["Productivity"].weight_applied) == pytest.approx(0.10)
    assert float(kpis["Productivity"].actual_value) == pytest.approx(AUGUST["Productivity"], abs=5e-5)
    assert float(kpis["Productivity"].target_value) == pytest.approx(0.8)
    assert float(kpis["Attendance"].target_value) == pytest.approx(0.65)
    payload_ratio = sum(float(item["contribution"]) for item in stored.record_payload["kpi_values"])
    assert payload_ratio == pytest.approx(AUGUST_RATIO, abs=1e-8)
    assert float(stored.score) == pytest.approx(round(AUGUST_RATIO * 100, 2), abs=0.001)
    assert float(stored.score) == pytest.approx(79.82, abs=0.001)
    old_four = (
        min(AUGUST["Attendance"] / 0.65, 1) * 0.70
        + min(AUGUST["Booking"] / 0.3, 1) * 0.10
        + min(AUGUST["Quality"] / 0.95, 1) * 0.10
        + min(AUGUST["Other"] / 0.75, 1) * 0.10
    )
    assert round(old_four * 100, 2) == 77.19
    assert float(stored.score) != pytest.approx(77.19, abs=0.01)
    thresholds = load_team_config("Outbound")["grade_thresholds"]
    assert stored.grade == calculate_grade(float(stored.score), thresholds)
    assert next(item["weight"] for item in load_team_config("Outbound")["kpis"] if item["key"] == "Attendance") == 0.70
    assert all(item["key"] != "Productivity" for item in load_team_config("Outbound")["kpis"])
    payload = stored.record_payload
    assert payload["actual"]["productivity_rate"] == pytest.approx(AUGUST["Productivity"])
    assert payload["raw_data"]["T.Attend%"] == pytest.approx(0.65)
    assert payload["raw_data"]["T.Attend%"] != "15:36:00"
    assert payload["calls"]["aht_raw"] == "00:02:30"
    assert payload["actual"]["booking_rate"] == pytest.approx(AUGUST["Booking"])
    config_keys = {row.kpi_key for row in db.query(TeamKPIConfig)}
    assert "Productivity" not in config_keys
    assert float(db.query(TeamKPIConfig).filter(TeamKPIConfig.kpi_key == "Attendance").one().weight) == pytest.approx(0.70)

    [resolved] = DashboardRecordService(db).resolve_records([stored])
    assert resolved.actual.productivity_rate == pytest.approx(AUGUST["Productivity"])
    assert resolved.calls.aht_raw == "00:02:30"
    assert float(resolved.evaluation.score) == pytest.approx(79.82, abs=0.001)
    productivity = next(item for item in resolved.kpi_values if item["kpi_key"] == "Productivity")
    assert productivity["weight_applied"] == pytest.approx(0.10)


def test_july_upload_keeps_the_four_kpi_score(db, monkeypatch):
    _bind(db, monkeypatch)
    blob = _workbook(
        datetime(2026, 7, 15), JULY, productivity=None, role="Emp", sheet="Outbound Offshore Call Center",
    )
    _seeder().process_uploaded_file("july.xlsx", blob, dry_run=False)
    stored = db.query(PerformanceRecord).one()
    kpis = _by_key(stored.kpi_values)
    assert set(kpis) == {"Booking", "Attendance", "Quality", "Other"}
    assert float(kpis["Attendance"].weight_applied) == pytest.approx(0.70)
    assert float(kpis["Attendance"].target_value) == pytest.approx(0.65)
    payload_ratio = sum(float(item["contribution"]) for item in stored.record_payload["kpi_values"])
    assert payload_ratio == pytest.approx(JULY_RATIO, abs=1e-8)
    assert float(stored.score) == pytest.approx(82.85, abs=0.001)
    assert stored.record_payload["actual"]["productivity_rate"] is None


def test_missing_august_productivity_is_refused_on_dry_run_and_commit(db, monkeypatch):
    _bind(db, monkeypatch)
    blob = _august_bytes(productivity=None)
    seeder = _seeder()
    for dry_run in (True, False):
        with pytest.raises(UploadProcessingError, match="Productivity"):
            seeder.process_uploaded_file("missing.xlsx", blob, dry_run=dry_run)
        db.rollback()
        assert db.query(PerformanceRecord).count() == 0


def test_september_does_not_inherit_august_and_refuses_an_approved_binding(db, monkeypatch):
    _bind(db, monkeypatch)
    actuals = {key: AUGUST[key] for key in ("Booking", "Attendance", "Other", "Quality")}
    blob = _workbook(datetime(2026, 9, 15), actuals, productivity=None, role="Emp", sheet="Outbound")
    _seeder().process_uploaded_file("september.xlsx", blob, dry_run=False)
    stored = db.query(PerformanceRecord).one()
    weights = {row.kpi_key: float(row.weight_applied) for row in stored.kpi_values}
    assert weights["Attendance"] == pytest.approx(0.70)
    assert "Productivity" not in weights
    assert float(stored.score) != pytest.approx(79.82, abs=0.01)
    _delete_record(db, stored)

    present = _workbook(
        datetime(2026, 9, 15), actuals, productivity=AUGUST["Productivity"], role="Emp", sheet="Outbound",
    )
    _seeder().process_uploaded_file("september-present.xlsx", present, dry_run=False)
    stored = db.query(PerformanceRecord).one()
    productivity = next(row for row in stored.kpi_values if row.kpi_key == "Productivity")
    assert float(productivity.weight_applied) == 0
    assert float(productivity.contribution) == 0
    assert float(productivity.actual_value) == pytest.approx(AUGUST["Productivity"], abs=5e-5)
    assert stored.record_payload["actual"]["productivity_rate"] == pytest.approx(AUGUST["Productivity"])
    _delete_record(db, stored)

    team = db.query(Team).filter(Team.name == "Outbound").one()
    db.add(_approved(team.id, 9, _lines(attendance_target=0.65)))
    db.commit()
    with pytest.raises(OutboundUnsupportedApprovedBinding):
        _seeder().process_uploaded_file("september-approved.xlsx", blob, dry_run=False)
    db.rollback()
    assert db.query(PerformanceRecord).count() == 0


def test_fixed_august_target_blocks_dry_run_and_commit_without_override(db, monkeypatch):
    _bind(db, monkeypatch)
    team = Team(id=uuid.uuid4(), name="Outbound", db_name="Outbound", display_name="Outbound", region="EGY", team_level="employee")
    db.add(team)
    db.flush()
    db.add(_approved(team.id, 8, _lines(attendance_target=0.55)))
    db.commit()
    blob = _august_bytes()
    seeder = _seeder()
    for dry_run in (True, False):
        with pytest.raises(TargetConflict) as caught:
            seeder.process_uploaded_file("conflict.xlsx", blob, dry_run=dry_run)
        conflict = next(item for item in caught.value.data["conflicts"] if item["kpi_key"] == "Attendance")
        assert conflict["workbook_target"] == pytest.approx(0.65)
        assert conflict["approved_target"] == pytest.approx(0.55)
        db.rollback()
        assert db.query(PerformanceRecord).count() == 0


def test_pinned_read_keeps_the_rich_payload_when_current_config_changes(monkeypatch):
    monkeypatch.setattr(
        "services.dashboard_record_service.load_team_config",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(AssertionError("current config must not be read")),
    )
    team = SimpleNamespace(name="Outbound", db_name="Outbound", display_name=None)
    employee = SimpleNamespace(employee_id="ANON-1", name="Anonymous Agent", team=team, region="EGY", position_name=None)
    payload = {
        "id": "payload-id",
        "employee_id": "PAYLOAD",
        "employee_name": "Payload Name",
        "team": "Outbound",
        "month": "August",
        "year": 2026,
        "status": "Below",
        "upload_id": "batch-august",
        "meta": {"source": "stored-payload"},
        "calls": {"inbound": 1, "outbound": 4, "total_handled": 9, "abandoned": 0, "aht_raw": "00:02:30"},
        "geo": {"bookings": {"dubai": 3, "sharjah": 0, "ajman": 0, "clinics": 0}, "attended": {"dubai": 1, "sharjah": 0, "ajman": 0, "clinics": 0}},
        "actual": {"productivity_rate": AUGUST["Productivity"], "booking_rate": AUGUST["Booking"], "attend_rate": AUGUST["Attendance"]},
        "achievement": {"productivity_ach": AUGUST["Productivity"] / 0.8},
        "evaluation": {
            "score": 1.0,
            "grade": "A",
            "manager_notes": "kept note",
            "root_cause": {"kpi": "Attendance", "impact_pct": 12.5, "actual": AUGUST["Attendance"], "target": 0.65},
            "suggested_action": "review attendance",
        },
        "raw_data": {"Productivity": AUGUST["Productivity"], "Available Time": 0.25},
        "kpi_values": [{
            "kpi_key": "Productivity",
            "label": "Payload label",
            "actual_value": 0.1,
            "target_value": 0.1,
            "achievement_ratio": 1,
            "weight_applied": 0.7,
            "contribution": 0.7,
        }],
        "evaluation_basis": {
            "pinned": True,
            "lines": [{"kpi_key": "Productivity", "label": "Productivity", "direction": "higher_better", "unit": "%", "weight": 0.10}],
        },
    }
    item = SimpleNamespace(
        id="sql-id",
        employee=employee,
        team=team,
        month="August",
        year=2026,
        region="EGY",
        branch_key=None,
        performance_level="Employee",
        position_name=None,
        status="Meets",
        score=79.82,
        grade="C",
        upload_id=uuid.uuid4(),
        record_payload=payload,
        kpi_values=[SimpleNamespace(
            kpi_key="Productivity",
            actual_value=AUGUST["Productivity"],
            target_value=0.8,
            achievement_ratio=AUGUST["Productivity"] / 0.8,
            weight_applied=0.10,
            contribution=0.09725868055555556,
        )],
    )
    [resolved] = DashboardRecordService(object()).resolve_records([item])
    assert resolved.employee_id == "ANON-1"
    assert resolved.status == "Meets"
    assert resolved.upload_id == "batch-august"
    assert resolved.meta == {"source": "stored-payload"}
    assert resolved.calls.outbound == 4
    assert resolved.geo.bookings.dubai == 3
    assert resolved.actual.productivity_rate == pytest.approx(AUGUST["Productivity"])
    assert resolved.actual.booking_rate == pytest.approx(AUGUST["Booking"])
    assert resolved.evaluation.score == pytest.approx(79.82)
    assert resolved.evaluation.grade == "C"
    assert resolved.evaluation.manager_notes == "kept note"
    assert resolved.evaluation.root_cause.kpi == "Attendance"
    assert resolved.evaluation.suggested_action == "review attendance"
    assert resolved.kpi_values[0]["weight_applied"] == pytest.approx(0.10)
    assert resolved.kpi_values[0]["actual_value"] == pytest.approx(AUGUST["Productivity"])
    assert resolved.kpi_values[0]["label"] == "Productivity"
    assert resolved.raw_data["Available Time"] == 0.25

    clear_serialization_cache()
    body = serialize_performance_record(resolved)
    assert body["upload_id"] == "batch-august"
    assert body["meta"]["source"] == "stored-payload"
    assert body["calls"]["outbound"] == 4
    assert body["actual"]["productivity_rate"] == pytest.approx(AUGUST["Productivity"])
    assert body["evaluation"]["manager_notes"] == "kept note"
    assert body["evaluation"]["root_cause"]["kpi"] == "Attendance"
    flat = ReportExporter.flatten_record(resolved)
    assert flat["Manager Notes"] == "kept note"
    assert flat["Root Cause"] == "Attendance"
    assert flat["Outbound Calls"] == 4
    assert flat["Productivity Actual"] == pytest.approx(AUGUST["Productivity"])
    rolled = aggregate_kpi_metric([item for item in resolved.kpi_values if item["kpi_key"] == "Productivity"])
    assert rolled.actual == pytest.approx(AUGUST["Productivity"])
    assert isinstance(resolved, SchemaRecord)
    assert resolved.evaluation.score != EvaluationData(score=1, grade="A").score


def _lines(attendance_target: float) -> list[dict]:
    weights = {"Booking": 0.10, "Attendance": 0.60, "Other": 0.10, "Quality": 0.10, "Productivity": 0.10}
    return [
        {
            "kpi_key": key,
            "label": key,
            "direction": "higher_better",
            "unit": "%",
            "weight": weight,
            "target": attendance_target if key == "Attendance" else TARGETS[key],
            "target_mode": "fixed",
        }
        for key, weight in weights.items()
    ]


def _approved(team_id, month: int, lines: list[dict]) -> TeamConfigurationVersion:
    """Exact monthly row: real month name, matching until endpoints, and published_at."""
    from datetime import timezone

    from services.evaluation.periods import month_name

    label = month_name(month)
    return TeamConfigurationVersion(
        id=uuid.uuid4(),
        team_id=team_id,
        version_number=1,
        status="approved",
        effective_month=label,
        effective_year=2026,
        config_snapshot={"lines": lines, "grade_thresholds": {"A": 95, "B": 85, "C": 75, "D": 65}},
        config_checksum="a" * 64,
        effective_from_month=month,
        effective_from_year=2026,
        effective_until_month=month,
        effective_until_year=2026,
        performance_level="Employee",
        position_name="",
        published_at=datetime.now(timezone.utc),
    )
