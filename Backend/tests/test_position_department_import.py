"""The management-template upload repairs the workbook, then splits by level.

Corporate rows are the only ones written to the two management-KPI stores.
Employee and Managerial rows become pinned department records. The fixture's
departmental block is one month early and overweight, so a June record at the
New Temp weights proves the shipped repair ran before the import.
"""
import io
from datetime import datetime

import pytest
from openpyxl import Workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from models.models import (
    Base,
    Employee,
    KPIValue,
    ManagementKPIConfig,
    ManagementKPISnapshot,
    PerformanceRecord,
    Team,
)
from services.dashboard_record_service import DashboardRecordService
from services.management_template_import import import_uploaded_management_template
from services.position_department_import_service import PositionDepartmentImportService
from utils.team_identity import create_management_team_identity

DATA_HEADERS = [
    "Employee ID", "Team", "Employee Name", "Position", "Performance Level", "Period",
    "Perspective", "KPI", "Direction", "Weight", "Target Value", "Target Unit",
    "Actual Value", "Actual Unit", "Score %",
]
CONFIG_HEADERS = [
    "Team", "Position", "Employee Override ID (Optional)", "Performance Level",
    "Perspective", "KPI", "Direction", "Weight", "Target Unit",
]


def _db():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    return sessionmaker(bind=engine)()


def _workbook() -> bytes:
    workbook = Workbook()
    data = workbook.active
    data.title = "KPI's Data"
    config = workbook.create_sheet("KPI Config")
    source = workbook.create_sheet("New Temp")
    data.append(DATA_HEADERS)
    source.append(DATA_HEADERS)
    config.append(CONFIG_HEADERS)

    may = datetime(2026, 5, 1)
    june = datetime(2026, 6, 1)
    # Corporate stays outside the New Temp suffix. Perspective is blank on the
    # data sheet and comes from KPI Config.
    data.append([
        "SGHD90000", "Finance", "Corporate Director", "CFO", "Corporate", june,
        None, "Revenue", "Higher is Better", 1, 100, "%", 100, "%", 1,
    ])
    # Departmental suffix: dates and weights are wrong. New Temp is the truth.
    departmental = [
        ["SGHD90001", "Finance", "Ada Accountant", "General Accountant", "Employee", may,
         None, "Collections", "Lower is Better", 0.9, 100, "%", 150, "%", 1],
        ["SGHD90001", "Finance", "Ada Accountant", "General Accountant", "Employee", may,
         None, "Accuracy", "Lower is Better", 0.9, 80, "%", 40, "%", 0.5],
        ["SGHD90002", "Call Center", "Sam Leader", "Inbound Team Leader", "Managerial", may,
         None, "Booking CR%", "Higher is Better", 0.5, 100, "%", 90, "%", 0.9],
    ]
    corrected = [
        ["SGHD90001", "Finance", "Ada Accountant", "General Accountant", "Employee", "Jun",
         None, "Collections", "Higher is Better", 0.6, 100, "%", 150, "%", 1],
        ["SGHD90001", "Finance", "Ada Accountant", "General Accountant", "Employee", "Jun",
         None, "Accuracy", "Higher is Better", 0.4, 80, "%", 40, "%", 0.5],
        ["SGHD90002", "Call Center", "Sam Leader", "Inbound Team Leader", "Managerial", "Jun",
         None, "Booking CR%", "Higher is Better", 1, 100, "%", 90, "%", 0.9],
    ]
    for row in departmental:
        data.append(row)
    for row in corrected:
        source.append(row)

    config.append(["Finance", "CFO", None, "Corporate", "Financial", "Revenue", "Higher is Better", 1, "%"])
    config.append(["Finance", "General Accountant", None, "Employee", "Financial", "Collections", "Higher is Better", 0.6, "%"])
    config.append(["Finance", "General Accountant", None, "Employee", "Internal Process", "Accuracy", "Higher is Better", 0.4, "%"])
    config.append(["Call Center", "Inbound Team Leader", None, "Managerial", "Customer", "Booking CR%", "Higher is Better", 1, "%"])

    payload = io.BytesIO()
    workbook.save(payload)
    return payload.getvalue()


def _seed_managerial_scorecard(session) -> None:
    team = create_management_team_identity(session, "Sales", region="EGY")
    session.add(ManagementKPIConfig(
        team_id=team.id,
        performance_level="Managerial",
        position_name="Legacy Director",
        employee_identifier=None,
        perspective_key="Financial",
        kpi_key="legacy_revenue",
        kpi_label="Legacy Revenue",
        direction="higher_better",
        weight=1,
        effective_month="January",
        effective_year=2026,
    ))
    session.add(ManagementKPISnapshot(
        team_id=team.id,
        employee_identifier="SGHD80000",
        employee_name="Legacy Director",
        position_name="Legacy Director",
        performance_level="Managerial",
        month="January",
        year=2026,
        perspective_key="Financial",
        kpi_key="legacy_revenue",
        kpi_label="Legacy Revenue",
        actual_value=10,
    ))
    session.commit()


def test_template_upload_routes_levels_and_pins_department_records():
    session = _db()
    try:
        _seed_managerial_scorecard(session)
        result = import_uploaded_management_template(
            session,
            _workbook(),
            updated_by="tester",
            source_filename="Template_Managment.xlsx",
        )

        assert result["position_records"] == 2
        assert result["position_teams"] == ["Call Center", "Finance"]
        assert result["levels"] == ["Corporate"]

        configs = session.query(ManagementKPIConfig).all()
        snapshots = session.query(ManagementKPISnapshot).all()
        assert configs and snapshots
        assert {row.performance_level for row in configs} == {"Corporate"}
        assert {row.performance_level for row in snapshots} == {"Corporate"}
        assert {row.perspective_key for row in configs} == {"Financial"}
        assert {row.kpi_label for row in snapshots} == {"Revenue"}

        finance_employee = session.query(Employee).filter(Employee.employee_id == "SGHD90001").one()
        finance_team = session.query(Team).filter(Team.id == finance_employee.team_id).one()
        assert finance_team.name == "Finance"
        assert finance_team.team_level == "employee"
        assert finance_team.region == "All"
        finance = session.query(PerformanceRecord).filter(PerformanceRecord.employee_id == finance_employee.id).one()
        assert finance.performance_level == "Employee"
        assert finance.position_name == "General Accountant"
        assert finance.month == "June"
        assert finance.year == 2026
        assert float(finance.score) == pytest.approx(80.0)
        assert finance.grade == "C"
        assert finance.record_payload["source_performance_level"] == "Employee"
        assert finance.record_payload["position"] == "General Accountant"
        assert finance.record_payload["role"] == "General Accountant"
        finance_kpis = {row.kpi_key: row for row in session.query(KPIValue).filter(KPIValue.record_id == finance.id)}
        collections = next(row for row in finance_kpis.values() if float(row.weight_applied) == pytest.approx(0.6))
        accuracy = next(row for row in finance_kpis.values() if float(row.weight_applied) == pytest.approx(0.4))
        assert float(collections.achievement_ratio) == pytest.approx(1.0)
        assert float(collections.actual_value) == pytest.approx(1.5)
        assert float(collections.target_value) == pytest.approx(1)
        assert float(collections.contribution) == pytest.approx(0.6)
        assert float(accuracy.achievement_ratio) == pytest.approx(0.5)
        assert float(accuracy.contribution) == pytest.approx(0.2)
        payload_by_key = {row["kpi_key"]: row for row in finance.record_payload["kpi_values"]}
        assert payload_by_key[collections.kpi_key]["evaluation_pinned"] is True
        assert payload_by_key[collections.kpi_key]["label"] == "Collections"
        assert payload_by_key[collections.kpi_key]["perspective"] == "Financial"
        assert payload_by_key[collections.kpi_key]["direction"] == "higher_better"
        assert payload_by_key[accuracy.kpi_key]["perspective"] == "Internal Process"
        assert payload_by_key[accuracy.kpi_key]["evaluation_pinned"] is True

        leader = session.query(Employee).filter(Employee.employee_id == "SGHD90002").one()
        call_center = session.query(Team).filter(Team.id == leader.team_id).one()
        assert call_center.name == "Call Center"
        assert call_center.region == "EGY"
        leader_record = session.query(PerformanceRecord).filter(PerformanceRecord.employee_id == leader.id).one()
        assert leader_record.performance_level == "Employee"
        assert leader_record.position_name == "Inbound Team Leader"
        assert leader_record.month == "June"
        assert float(leader_record.score) == pytest.approx(90.0)
        assert leader_record.record_payload["source_performance_level"] == "Managerial"
        booking = session.query(KPIValue).filter(KPIValue.record_id == leader_record.id).one()
        assert float(booking.weight_applied) == pytest.approx(1.0)
        assert float(booking.achievement_ratio) == pytest.approx(0.9)
        assert float(booking.contribution) == pytest.approx(0.9)
        assert leader_record.record_payload["kpi_values"][0]["evaluation_pinned"] is True

        finance_view = DashboardRecordService(session).list_records(team="Finance")
        assert len(finance_view) == 1
        assert finance_view[0].position == "General Accountant"
        assert finance_view[0].performance_level == "Employee"
        viewed = {row["label"]: row for row in finance_view[0].kpi_values}
        assert viewed["Collections"]["evaluation_pinned"] is True
        assert viewed["Collections"]["achievement_ratio"] == pytest.approx(1.0)
        assert viewed["Collections"]["contribution"] == pytest.approx(0.6)
        assert viewed["Accuracy"]["unit"] == "%"
        assert viewed["Accuracy"]["direction"] == "higher_better"

        call_center_view = DashboardRecordService(session).list_records(team="Call Center")
        assert any(row.employee_id == "SGHD90002" and row.position == "Inbound Team Leader" for row in call_center_view)
    finally:
        session.close()


def test_existing_call_center_team_is_corrected_to_egypt():
    session = _db()
    session.add(Team(
        name="Call Center",
        db_name="call_center",
        display_name="Call Center",
        region="UAE",
        team_level="employee",
        is_active=True,
    ))
    session.commit()

    team = PositionDepartmentImportService(session)._employee_team("Call Center")

    assert team.region == "EGY"
    session.close()


def _rate_workbook() -> bytes:
    workbook = Workbook()
    data = workbook.active
    data.title = "KPI's Data"
    config = workbook.create_sheet("KPI Config")
    source = workbook.create_sheet("New Temp")
    data.append(DATA_HEADERS)
    source.append(DATA_HEADERS)
    config.append(CONFIG_HEADERS)
    september = datetime(2026, 9, 1)
    august = datetime(2026, 8, 1)
    stored = [
        ["SGHD91011", "Human Resources", "Hana Specialist", "HR Specialist", "Employee", september,
         "Internal Process", "Onboarding Rate", "Higher is Better", 1, 54, "%", 43, "%", 43 / 54],
        ["SGHD91012", "Human Resources", "Omar Supervisor", "HR Operation Supervisor", "Employee", september,
         "Internal Process", "File Completion", "Higher is Better", 1, 107, "%", 95, "%", 95 / 107],
        ["SGHD91013", "Marketing", "Mona Buyer", "Media Buyer & SEO Manager", "Employee", september,
         "Financial", "Budget Compliance", "Higher is Better", 1, 997017.05, "%", 997333.49, "%", 1],
        ["SGHD91014", "IT", "Gamal Support", "IT Technical Support", "Managerial", august,
         "Internal Process", "Average Ticket Resolution Time", "Lower is Better", 1, 1, "%", 0.9, "%", 0.9],
        ["SGHD91014", "IT", "Gamal Support", "IT Technical Support", "Managerial", september,
         "Internal Process", "Average Ticket Resolution Time", "Higher Better", 1, 48, "%", 24, "%", 1],
    ]
    source_months = ["Sep", "Sep", "Sep", "Aug", "Sep"]
    for row, month_name in zip(stored, source_months):
        data.append(row)
        source_row = list(row)
        source_row[5] = month_name
        source.append(source_row)
    payload = io.BytesIO()
    workbook.save(payload)
    return payload.getvalue()


def test_position_import_stores_percent_rates_and_ticket_compliance():
    session = _db()
    try:
        import_uploaded_management_template(
            session,
            _rate_workbook(),
            updated_by="tester",
            source_filename="Template_Managment.xlsx",
        )
        records = session.query(PerformanceRecord).all()
        none_units = 0
        by_label = {}
        for record in records:
            for kpi in record.record_payload["kpi_values"]:
                if str(kpi.get("unit") or "").strip().casefold() == "none":
                    none_units += 1
                by_label[(record.record_payload["employee_id"], record.month, kpi["label"])] = kpi
        assert none_units == 0
        specialist = by_label[("SGHD91011", "September", "Onboarding Rate")]
        assert float(specialist["actual_value"]) == pytest.approx(43 / 54, abs=0.0001)
        assert float(specialist["target_value"]) == pytest.approx(1)
        assert specialist["unit"] == "%"
        supervisor = by_label[("SGHD91012", "September", "File Completion")]
        assert float(supervisor["actual_value"]) == pytest.approx(95 / 107, abs=0.0001)
        assert float(supervisor["target_value"]) == pytest.approx(1)
        budget = by_label[("SGHD91013", "September", "Budget Compliance")]
        assert float(budget["actual_value"]) == pytest.approx(997333.49 / 997017.05, abs=0.0001)
        assert float(budget["target_value"]) == pytest.approx(1)
        assert budget["unit"] == "%"
        august = by_label[("SGHD91014", "August", "Average Ticket Resolution Time")]
        september = by_label[("SGHD91014", "September", "Average Ticket Resolution Time")]
        assert august["unit"] == "%"
        assert august["direction"] == "higher_better"
        assert float(august["actual_value"]) == pytest.approx(0.9)
        assert float(august["target_value"]) == pytest.approx(1)
        assert september["unit"] == "%"
        assert september["direction"] == "higher_better"
        assert float(september["actual_value"]) == pytest.approx(1)
        assert float(september["target_value"]) == pytest.approx(1)
        assert float(september["achievement_ratio"]) == pytest.approx(1)
        assert float(august["achievement_ratio"]) == pytest.approx(0.9)
    finally:
        session.close()
