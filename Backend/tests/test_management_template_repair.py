"""The management-template repair is the shipped transform, exercised on a fixture and the real workbook."""
from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from services.management_template_repair import (
    CONFIG_SHEET,
    DATA_SHEET,
    HIGHER,
    LOWER,
    SOURCE_SHEET,
    repair_workbook,
)

REAL_WORKBOOK = Path(r"D:\Trend\Template_Managment.xlsx")
HEADERS = [
    "Employee ID", "Team", "Employee Name", "Position", "Performance Level", "Period",
    "Perspective", "KPI", "Direction", "Weight", "Target Value", "Target Unit",
    "Actual Value", "Actual Unit", "Score %", "Weighted Score %", "Status",
]
CONFIG_HEADERS = [
    "Team", "Position", "Employee Override ID (Optional)", "Performance Level",
    "Perspective", "KPI", "Direction", "Weight", "Target Unit",
]


def _achievement(direction: str, actual: float, target: float) -> float:
    if "lower" in direction.casefold():
        return 0.0 if actual == 0 else min(1.0, target / actual)
    return 0.0 if target == 0 else min(1.0, actual / target)


def _write_fixture(path: Path) -> None:
    book = Workbook()
    data = book.active
    data.title = DATA_SHEET
    data.append(HEADERS)
    config = book.create_sheet(CONFIG_SHEET)
    config.append(CONFIG_HEADERS)
    source = book.create_sheet(SOURCE_SHEET)
    source.append(HEADERS)

    # Incomplete duplicate of the later IT block, on a month New Temp does not contain.
    for kpi, weight in (("Ticket Resolution Rate", 0.35), ("User Satisfaction Score", 0.35)):
        data.append(_data_row(
            "SGHD70086", "IT", "Mohamed Gamal", "IT Technical Support", "Managerial",
            date(2026, 1, 1), kpi, "Higher Better", weight, 1, 1, 1,
        ))

    workforce = [
        ("Forecast Accuracy", 0.25, 0.30),
        ("Schedule Adherence", 0.25, 0.25),
        ("Schedule Adherence Rate", 0.20, 0.20),
        ("Occupancy", 0.15, 0.15),
        ("Service Level", 0.15, 0.30),
    ]
    for kpi, clean, collided in workforce:
        source.append(_data_row("SGHD70094", " Call Center ", "Mahmoud Fathy Mohamed", "Workforce analyst", "Employee", "Jun", kpi, HIGHER, clean, 1, 1, 1))
        data.append(_data_row("SGHD70094", " Call Center ", "Mahmoud Fathy Mohamed", "Workforce analyst", "Employee", date(2026, 5, 1), kpi, "Higher Better", collided, 1, 1, 1))

    for code in ("SGHD70001", "SGHD70005", "SGHD70002", "SGHD70016"):
        leader_rows = [
            ("Team Adherence", 0.35, 0.35, "Higher Better", 1, 1, 1),
            ("Team Shrinkage", 0.35, 0.35, "Higher Better", 0.07, 0.06, 1),
            ("Team Attendance", 0.15, 0.10, "Higher Better", 1, 1, 1),
            ("Team Quality", 0.15, 0.15, "Higher Better", 1, 1, 1),
        ]
        for kpi, clean, collided, direction, target, actual, score in leader_rows:
            source.append(_data_row(code, "Call Center", "Team Leader", "Inbound Team Leader", "Managerial", "Jun", kpi, HIGHER, clean, target, actual, score))
            data.append(_data_row(code, "Call Center", "Team Leader", " Inbound Team Leader ", "Managerial", date(2026, 5, 1), kpi, direction, collided, target, actual, score))

    for month_name, shifted, target, actual, score in (
        ("Aug", date(2026, 7, 31), 0.85, 0.80, 0.9411764705882354),
        ("Sep", date(2026, 8, 31), 0.85, 0.82, 0.9647058823529412),
    ):
        source.append(_data_row("SGHD70227", "Offshore Admin & Management", "Anwaar", "Offshore PMO", "Employee", month_name, "Schedule Variance", "Lower Better", 0.30, target, actual, score))
        source.append(_data_row("SGHD70227", "Offshore Admin & Management", "Anwaar", "Offshore PMO", "Employee", month_name, "Portfolio Completion Rate", HIGHER, 0.70, 1, 1, 1))
        data.append(_data_row("SGHD70227", " Offshore Admin & Management ", "Anwaar", "Offshore PMO", "Employee", shifted, "Schedule Variance", "Lower Better", 0.30, target, actual, score))
        data.append(_data_row("SGHD70227", "Offshore Admin & Management", "Anwaar", "Offshore PMO", "Employee", shifted, "Portfolio Completion Rate", "Higher Better", 0.70, 1, 1, 1))

    for month_name, shifted, ticket_direction, target, actual, score in (
        ("Aug", date(2026, 7, 31), "Lower Better", 1, 0.9, 0.9),
        ("Sep", date(2026, 8, 31), "Higher Better", 48, 24, 1),
    ):
        for kpi, direction, weight, kpi_target, kpi_actual, kpi_score in (
            ("Ticket Resolution Rate", HIGHER, 0.35, 1, 1, 1),
            ("User Satisfaction Score", HIGHER, 0.35, 1, 1, 1),
            ("Average Ticket Resolution Time", ticket_direction, 0.30, target, actual, score),
        ):
            source.append(_data_row("SGHD70086", "IT", "Mohamed Gamal", "IT Technical Support", "Managerial", month_name, kpi, direction, weight, kpi_target, kpi_actual, kpi_score))
            data.append(_data_row("SGHD70086", "IT", "Mohamed Gamal", "IT Technical Support", "Managerial", shifted, kpi, direction if kpi != "Average Ticket Resolution Time" else ticket_direction, weight, kpi_target, kpi_actual, kpi_score))

    config.append(["Call Center", "Workforce analyst", None, "Managerial", "Customer", "Forecast Accuracy", "Higher Better", 0.20, "%"])
    config.append(["Call Center", "Workforce analyst", None, "Employee", "Customer", "Forecast Accuracy", "Higher Better", 0.30, "%"])
    config.append(["Call Center", " Inbound Team Leader", None, "Managerial", "Internal Process", "Team Shrinkage", "Higher Better", 0.35, "%"])
    config.append(["Call Center", "Inbound Team Leader", None, "Managerial", "Internal Process", "Team Shrinkage", "Higher Better", 0.35, "%"])
    config.append(["Offshore Admin & Management", "Offshore PMO", None, "Employee", "Internal Process", "Schedule Variance", "Lower Better", 0.30, "%"])
    config.append(["IT", "IT Technical Support", None, "Managerial", "Internal Process", "Average Ticket Resolution Time", "Lower Better", 0.30, "%"])
    book.save(path)


def _data_row(code, team, name, position, level, period, kpi, direction, weight, target, actual, score):
    return [code, team, name, position, level, period, "Internal Process", kpi, direction, weight, target, "%", actual, "%", score, None, "Ready"]


def _rows(sheet) -> list[tuple]:
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    return [row for row in rows if any(cell is not None and str(cell).strip() for cell in row)]


def _period(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def assert_criterion_one(path: Path) -> None:
    book = load_workbook(path, data_only=False)
    data = _rows(book[DATA_SHEET])
    config = _rows(book[CONFIG_SHEET])

    stray = [row for row in data if str(row[0]) == "SGHD70086" and _period(row[5]) == date(2026, 1, 1)]
    assert stray == []

    workforce = [row for row in data if str(row[0]) == "SGHD70094"]
    assert [row[9] for row in workforce] == [0.25, 0.25, 0.20, 0.15, 0.15]
    assert { _period(row[5]) for row in workforce } == {date(2026, 6, 1)}
    assert all(str(row[1]) == "Call Center" for row in workforce)

    for code in ("SGHD70001", "SGHD70005", "SGHD70002", "SGHD70016"):
        attendance = [row for row in data if str(row[0]) == code and row[7] == "Team Attendance"]
        assert len(attendance) == 1
        assert attendance[0][9] == 0.15
        assert _period(attendance[0][5]) == date(2026, 6, 1)
        shrinkage = [row for row in data if str(row[0]) == code and row[7] == "Team Shrinkage"]
        assert shrinkage[0][8] == LOWER
        assert _achievement(shrinkage[0][8], shrinkage[0][12], shrinkage[0][10]) == pytest.approx(1.0)

    schedule = [row for row in data if row[7] == "Schedule Variance"]
    assert schedule
    assert all(row[8] == HIGHER for row in schedule)
    august = next(row for row in schedule if _period(row[5]) == date(2026, 8, 1))
    september = next(row for row in schedule if _period(row[5]) == date(2026, 9, 1))
    assert _achievement(august[8], august[12], august[10]) == pytest.approx(0.941, abs=0.001)
    assert _achievement(september[8], september[12], september[10]) == pytest.approx(0.965, abs=0.001)

    tickets = [row for row in data if row[7] == "Average Ticket Resolution Time"]
    august_ticket = next(row for row in tickets if _period(row[5]) == date(2026, 8, 1))
    september_ticket = next(row for row in tickets if _period(row[5]) == date(2026, 9, 1))
    assert _achievement(august_ticket[8], august_ticket[12], august_ticket[10]) == pytest.approx(0.90)
    assert _achievement(september_ticket[8], september_ticket[12], september_ticket[10]) == pytest.approx(1.0)
    assert august_ticket[8] == HIGHER
    assert september_ticket[8] == LOWER

    totals: dict[tuple, float] = defaultdict(float)
    for row in data:
        totals[(row[0], row[3], _period(row[5]))] += float(row[9])
    assert totals
    for key, total in totals.items():
        assert total == pytest.approx(1.0, abs=0.011), key

    for row in data:
        for index in (1, 2, 3, 4, 7, 8):
            if isinstance(row[index], str):
                assert row[index] == row[index].strip()

    shrinkage_config = [row for row in config if row[5] == "Team Shrinkage"]
    assert len(shrinkage_config) == 1
    assert shrinkage_config[0][6] == LOWER
    schedule_config = [row for row in config if row[5] == "Schedule Variance"]
    assert schedule_config[0][6] == HIGHER
    # Both Workforce analyst levels stay; lookup scope is the level, not the KPI name alone.
    forecast = [row for row in config if row[5] == "Forecast Accuracy"]
    assert {row[3] for row in forecast} == {"Employee", "Managerial"}
    book.close()


def test_repair_fixes_the_five_workbook_bugs(tmp_path: Path):
    path = tmp_path / "template.xlsx"
    _write_fixture(path)
    repair_workbook(path)
    assert_criterion_one(path)
    repair_workbook(path)
    assert_criterion_one(path)


def test_real_management_workbook_matches_new_temp_after_repair():
    if not REAL_WORKBOOK.exists():
        pytest.skip("management workbook is not on this machine")
    repair_workbook(REAL_WORKBOOK)
    book = load_workbook(REAL_WORKBOOK, data_only=False)
    values = load_workbook(REAL_WORKBOOK, data_only=True)
    data = book[DATA_SHEET]
    source = book[SOURCE_SHEET]
    value_data = values[DATA_SHEET]
    totals: dict[tuple, float] = defaultdict(float)
    whitespace = 0
    for number in range(2, data.max_row + 1):
        code = data.cell(number, 1).value
        if code is None:
            continue
        for column in (1, 2, 3, 4, 5, 7, 8, 9, 12, 14):
            cell = data.cell(number, column).value
            if isinstance(cell, str) and cell != cell.strip():
                whitespace += 1
        weight = data.cell(number, 10).value
        if isinstance(weight, str) and weight.strip().lower() == "view":
            weight = 0.0001
        totals[(str(code).strip(), str(data.cell(number, 4).value).strip(), _period(data.cell(number, 6).value))] += float(weight)
    assert whitespace == 0
    for key, total in totals.items():
        assert total == pytest.approx(1.0, abs=0.011), (key, total)

    tickets = {}
    schedule = []
    for number in range(2, value_data.max_row + 1):
        kpi = str(value_data.cell(number, 8).value or "")
        period = _period(value_data.cell(number, 6).value)
        actual = value_data.cell(number, 13).value
        target = value_data.cell(number, 11).value
        direction = str(data.cell(number, 9).value or "")
        if kpi == "Average Ticket Resolution Time" and period in (date(2026, 8, 1), date(2026, 9, 1)):
            tickets[period] = _achievement(direction, float(actual), float(target))
        if kpi == "Schedule Variance":
            schedule.append(direction)
    assert tickets[date(2026, 8, 1)] == pytest.approx(0.90)
    assert tickets[date(2026, 9, 1)] == pytest.approx(1.0)
    assert schedule and all(item == HIGHER for item in schedule)

    source_keys = []
    for number in range(2, source.max_row + 1):
        code = source.cell(number, 1).value
        kpi = source.cell(number, 8).value
        if code is None and kpi is None:
            continue
        text = str(code).strip()
        if text.endswith(".0"):
            text = text[:-2]
        if text.isdigit():
            text = f"SGHD{text}"
        source_keys.append((text.upper(), str(kpi).strip().casefold(), str(source.cell(number, 6).value).strip().casefold()[:3]))
    data_keys = []
    for number in range(2, data.max_row + 1):
        code = data.cell(number, 1).value
        kpi = data.cell(number, 8).value
        if code is None:
            continue
        period = _period(data.cell(number, 6).value)
        month = {6: "jun", 7: "jul", 8: "aug", 9: "sep", 5: "may", 4: "apr", 3: "mar"}.get(period.month if period else 0, "")
        data_keys.append((str(code).strip().upper(), str(kpi).strip().casefold(), month))
    assert source_keys == data_keys[-len(source_keys):]
    book.close()
    values.close()
