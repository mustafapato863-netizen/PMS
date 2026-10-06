"""PR #15 QA findings on e0d1c77 (/workspace/qa/pms-pr15/BUGS.md).

BUG-1  read-time corrected records keep a stale status ("Exceeds").
BUG-2  an unchanged, on-target KPI was labelled "on target but worsening".
BUG-3  Inbound "Other" Abandon Rate and Utilization rows averaged together.
"""
from __future__ import annotations

import io
import pytest
from pptx import Presentation

from models.schemas import EvaluationData, PerformanceRecord
from exports.insights_pptx_builder import build_insights_pptx
from services.insights_report_service import build_insights_snapshot
from services.insights_service import (
    InsightsService,
    _analysis_narrative,
    _directional_fields,
    _movement_positive,
)
from services.planning_service import PlanningService
from utils.kpi_direction import kpi_group_key, kpi_variant
from utils.performance_status import reconciled_status, status_for_grade


# ----------------------------------------------------------------- helpers

class _Repo:
    def __init__(self, records):
        self.records = records

    def get_all(self):
        return self.records


def _service(records):
    repository = _Repo(records)
    service = InsightsService(repository, PlanningService(repository))
    service._authorized_records = lambda _scope: (records, 0)
    return service


def _scope():
    return {
        "role": "Admin",
        "has_unrestricted_team_access": True,
        "legacy_unscoped": False,
        "accessible_teams": [],
        "accessible_team_levels": [],
    }


def _kpi(key, label, actual, target, weight, direction, *, unit="%"):
    if direction == "lower_better":
        ratio = 1.0 if actual <= 0 else min(target / actual, 1.0)
    else:
        ratio = min(actual / target, 1.0)
    return {
        "kpi_key": key, "label": label, "unit": unit, "direction": direction,
        "actual_value": actual, "target_value": target, "weight_applied": weight,
        "achievement_ratio": ratio, "contribution": ratio * weight,
    }


def _record(month, kpis, *, employee_id="E1", team="Submission", position="Submission Officer", score=100.0):
    return PerformanceRecord(
        id=f"{employee_id}_{team}_2026_{month}", employee_id=employee_id, employee_name=employee_id, team=team,
        month=month, year=2026, region="UAE", position=position, performance_level="Employee", status="Exceeds",
        evaluation=EvaluationData(score=score, grade="A"), kpi_values=kpis,
    )


def _analysis(workspace, key, label=None):
    return next(
        item for item in workspace.team_analyses
        if item.kpi_key == key and (label is None or item.title.startswith(label))
    )


IRR = "initial_rejection_rate"        # Submission, lower_better, target 0.05
SWD = "submission_within_due_date"    # Submission, higher_better, target 0.90


def _submission(month, rejection, on_time):
    return _record(month, [
        _kpi(IRR, "Initial Rejection Rate", rejection, .05, .6, "lower_better"),
        _kpi(SWD, "Submission Within Due Date", on_time, .9, .4, "higher_better"),
    ])


# ------------------------------------------------- BUG-2: movement helper

@pytest.mark.parametrize("direction", ["higher_better", "lower_better"])
def test_movement_is_none_for_equal_or_missing_previous(direction):
    assert _movement_positive(1.0, 1.0, direction) is None
    assert _movement_positive(.93, .93, direction) is None
    # Float noise below the change_value rounding is stable too.
    assert _movement_positive(.93, .93 + 1e-7, direction) is None
    assert _movement_positive(1.0, None, direction) is None
    assert _movement_positive(None, 1.0, direction) is None


def test_movement_still_reports_real_moves_in_both_directions():
    assert _movement_positive(.95, .90, "higher_better") is True
    assert _movement_positive(.90, .95, "higher_better") is False
    assert _movement_positive(.02, .04, "lower_better") is True
    assert _movement_positive(.04, .02, "lower_better") is False
    # 0.0001 is a visible move at 4-decimal precision.
    assert _movement_positive(.9301, .93, "higher_better") is True


@pytest.mark.parametrize(
    ("actual", "target", "direction"),
    [(1.0, 1.0, "higher_better"), (.03, .05, "lower_better"), (.80, .90, "higher_better")],
)
def test_narrative_for_a_flat_kpi_says_stable(actual, target, direction):
    text, movement, _missed, _met = _analysis_narrative("KPI", actual, actual, target, "%", direction)
    assert movement is None
    assert "remained stable" in text
    assert "declined" not in text and "improved" not in text
    assert _directional_fields(actual, actual, target, direction)["trend_status"] == "stable"


# ------------------------------------- BUG-2: Insights workspace (repro shape)

def test_flat_capped_kpis_at_100_percent_are_on_target_not_worsening_for_both_directions():
    # Marketing COW-02 shape: every KPI at 100% achievement, unchanged month on month.
    workspace = _service([
        _submission("May", .03, .93),
        _submission("June", .03, .93),
    ]).generate_workspace(_scope(), month="June", year=2026)

    for key, label in ((IRR, "Initial Rejection Rate"), (SWD, "Submission Within Due Date")):
        item = _analysis(workspace, key)
        assert item.detail.change_value == 0
        assert item.detail.trend_status == "stable"
        assert item.detail.target_status == "met"
        assert "worsening" not in item.title.casefold()
        assert item.title in {f"{label} is on target", f"{label} is a positive score driver"}
        assert item.trend_label == "Target achieved"
        assert item.severity == "opportunity"
        assert "toward its" not in item.detail.recommended_focus
        assert item.detail.recommended_focus.startswith("Maintain")
        assert "remained stable" in item.explanation


def test_flat_exactly_at_target_is_on_target_for_both_directions():
    workspace = _service([
        _submission("May", .05, .90),
        _submission("June", .05, .90),
    ]).generate_workspace(_scope(), month="June", year=2026)
    for key in (IRR, SWD):
        item = _analysis(workspace, key)
        assert item.detail.trend_status == "stable"
        assert "worsening" not in item.title.casefold()
        assert item.trend_label == "Target achieved"


def test_flat_below_target_is_a_gap_without_improving_or_worsening_wording():
    workspace = _service([
        _submission("May", .08, .80),
        _submission("June", .08, .80),
    ]).generate_workspace(_scope(), month="June", year=2026)
    for key, label in ((IRR, "Initial Rejection Rate"), (SWD, "Submission Within Due Date")):
        item = _analysis(workspace, key)
        assert item.detail.trend_status == "stable"
        assert item.detail.target_status == "missed"
        assert item.title == f"{label} contributed to the performance gap"
        assert "improving" not in item.title.casefold() and "worsening" not in item.title.casefold()
        assert "remained stable" in item.explanation


def test_real_decline_within_target_is_still_a_watch_item():
    workspace = _service([
        _submission("May", .02, .99),
        _submission("June", .04, .95),
    ]).generate_workspace(_scope(), month="June", year=2026)
    rejection = _analysis(workspace, IRR)
    assert rejection.title == "Initial Rejection Rate is on target but worsening"
    assert rejection.severity == "information"
    assert rejection.trend_label == "Target achieved · Worsening"
    assert "rising toward its maximum allowed" in rejection.detail.recommended_focus
    on_time = _analysis(workspace, SWD)
    assert on_time.title == "Submission Within Due Date is on target but worsening"
    assert "falling toward its target" in on_time.detail.recommended_focus


def test_real_improvement_below_target_is_still_labelled_improving():
    workspace = _service([
        _submission("May", .10, .70),
        _submission("June", .08, .80),
    ]).generate_workspace(_scope(), month="June", year=2026)
    assert _analysis(workspace, IRR).title == "Initial Rejection Rate is improving but remains above target"
    assert _analysis(workspace, SWD).title == "Submission Within Due Date is improving but remains below target"


# ------------------------------------------- BUG-3: Inbound Other variants

def _inbound(employee_id, month, other):
    return _record(month, [
        _kpi("Attendance", "Attendance Rate", .75, .75, .7, "higher_better"),
        other,
    ], employee_id=employee_id, team="Inbound", position="Agent")


def _abandon(actual):
    return _kpi("Other", "Abandon Rate", actual, .01, .1, "lower_better")


def _utilization(actual):
    return _kpi("Other", "Utilization", actual, .85, .1, "higher_better")


def test_variant_group_key_separates_utilization_from_abandon_rate_only():
    assert kpi_variant({"kpi_key": "Other", "label": "Utilization"}) == "utilization"
    assert kpi_variant({"kpi_key": "Other", "label": "Abandon Rate"}) == ""
    assert kpi_variant({"kpi_key": "Other", "label": "Reachability"}) == ""
    assert kpi_group_key({"kpi_key": "Other", "label": "Utilization"}) != kpi_group_key(
        {"kpi_key": "Other", "label": "Abandon Rate"})
    # Ordinary KPIs still group by key alone (a renamed label does not split history).
    assert kpi_group_key({"kpi_key": "AHT", "label": "AHT"}) == kpi_group_key({"kpi_key": "AHT", "label": "Handle Time"})


def test_mixed_inbound_scope_never_averages_abandon_rate_with_utilization_in_insights():
    # QA repro: June has no-UTZ records (Abandon Rate) plus one real-UTZ record.
    records = [
        _inbound("EMP-01", "June", _abandon(.002)),
        _inbound("EMP-03", "June", _abandon(.008)),
        _inbound("EMP-05", "June", _utilization(.80)),
    ]
    workspace = _service(records).generate_workspace(_scope(), month="June", year=2026)
    others = [item for item in workspace.team_analyses if item.kpi_key == "Other"]
    by_direction = {item.detail.direction: item for item in others}

    assert len(others) == 2
    abandon, utilization = by_direction["lower_better"], by_direction["higher_better"]
    assert abandon.title.startswith("Abandon Rate")
    assert abandon.detail.current_value == pytest.approx(.005)
    assert abandon.detail.target_value == pytest.approx(.01)
    assert utilization.title.startswith("Utilization")
    assert utilization.detail.current_value == pytest.approx(.80)
    assert utilization.detail.target_value == pytest.approx(.85)


def test_mixed_inbound_scope_keeps_variants_apart_in_snapshot_and_pptx():
    def row(employee_id, other):
        return {
            "employee_id": employee_id, "employee_name": employee_id, "team": "Inbound", "position": "Agent",
            "performance_level": "Employee", "year": 2026, "month": "June", "score": 90.0,
            "kpis": [other],
        }

    records = [
        row("EMP-01", _abandon(.002)),
        row("EMP-03", _abandon(.008)),
        row("EMP-05", _utilization(.80)),
    ]
    report = {
        "period_label": "June 2026", "scope_label": "Inbound", "filters": {},
        "selected_records": records,
        "history": [{"key": "2026-06", "label": "June 2026", "records": records}],
        "actions": [],
    }
    snapshot = build_insights_snapshot(report)
    others = {kpi["label"]: kpi for kpi in snapshot["kpis"] if kpi["key"] == "Other"}
    assert set(others) == {"Abandon Rate", "Utilization"}
    assert others["Abandon Rate"]["direction"] == "lower_better"
    assert others["Abandon Rate"]["actual"] == pytest.approx(.005)
    assert others["Utilization"]["direction"] == "higher_better"
    assert others["Utilization"]["actual"] == pytest.approx(.80)
    assert others["Utilization"]["target"] == pytest.approx(.85)

    emp05 = next(person for person in snapshot["people"] if person["employee_id"] == "EMP-05")
    assert [kpi["label"] for kpi in emp05["kpis"]] == ["Utilization"]
    assert emp05["kpis"][0]["direction"] == "higher_better"

    deck = build_insights_pptx("June 2026", report)
    text = "\n".join(
        shape.text for slide in Presentation(io.BytesIO(deck)).slides for shape in slide.shapes if hasattr(shape, "text")
    )
    assert "Utilization" in text


# ----------------------------------------------------------- BUG-1: status

def test_status_rule_matches_every_import_path():
    assert [status_for_grade(g) for g in ("A", "B", "C", "D", "E", None)] == [
        "Exceeds", "Meets", "Meets", "Below", "Below", "Below"]
    # Recomputed only when the grade changed and the stored status was grade-derived.
    assert reconciled_status("Exceeds", "A", "B") == "Meets"
    assert reconciled_status("Exceeds", "A", "A") == "Exceeds"
    assert reconciled_status(None, "A", "B") == "Meets"
    assert reconciled_status("On Leave", "A", "B") == "On Leave"


# Record-level BUG-1 tests (real import -> DB -> dashboard path) live in
# tests/test_kpi_direction_followups.py next to the cw_error_free fixtures.
