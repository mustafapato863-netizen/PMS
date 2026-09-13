from io import BytesIO
from pathlib import Path, PurePosixPath
from zipfile import BadZipFile, ZipFile

from fastapi import HTTPException, UploadFile
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from config import settings


_XLSX_SIGNATURES = (b"PK\x03\x04", b"PK\x05\x06", b"PK\x07\x08")
_XLS_SIGNATURE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
_MAX_ARCHIVE_MEMBERS = 4096


def _validate_xlsx_structure(contents: bytes) -> None:
    """Reject malformed or resource-heavy OOXML workbooks before parsing them."""

    try:
        with ZipFile(BytesIO(contents)) as archive:
            members = archive.infolist()
            if len(members) > _MAX_ARCHIVE_MEMBERS:
                raise HTTPException(
                    status_code=413,
                    detail="The workbook contains too many internal files.",
                )

            total_uncompressed = 0
            for member in members:
                member_path = PurePosixPath(member.filename)
                if member_path.is_absolute() or ".." in member_path.parts:
                    raise HTTPException(
                        status_code=400,
                        detail="The uploaded workbook contains an invalid internal path.",
                    )
                total_uncompressed += max(0, member.file_size)
                if total_uncompressed > settings.MAX_UPLOAD_ARCHIVE_BYTES:
                    limit_mb = settings.MAX_UPLOAD_ARCHIVE_BYTES / (1024 * 1024)
                    raise HTTPException(
                        status_code=413,
                        detail=f"The workbook expands beyond the {limit_mb:g} MB processing limit.",
                    )
    except HTTPException:
        raise
    except (BadZipFile, OSError, ValueError) as exc:
        raise HTTPException(status_code=400, detail="The uploaded file is not a valid Excel workbook.") from exc

    workbook = None
    try:
        workbook = load_workbook(
            BytesIO(contents),
            read_only=True,
            data_only=True,
            keep_links=False,
        )
        if not workbook.sheetnames or len(workbook.sheetnames) > settings.MAX_UPLOAD_SHEETS:
            raise HTTPException(
                status_code=413,
                detail=f"The workbook must contain between 1 and {settings.MAX_UPLOAD_SHEETS} sheets.",
            )
        for worksheet in workbook.worksheets:
            if worksheet.max_row > settings.MAX_UPLOAD_ROWS_PER_SHEET:
                raise HTTPException(
                    status_code=413,
                    detail=f"Sheet '{worksheet.title}' exceeds the row processing limit.",
                )
            if worksheet.max_column > settings.MAX_UPLOAD_COLUMNS_PER_SHEET:
                raise HTTPException(
                    status_code=413,
                    detail=f"Sheet '{worksheet.title}' exceeds the column processing limit.",
                )
    except HTTPException:
        raise
    except (InvalidFileException, BadZipFile, OSError, ValueError, KeyError) as exc:
        raise HTTPException(status_code=400, detail="The uploaded file is not a valid Excel workbook.") from exc
    finally:
        if workbook is not None:
            workbook.close()


async def read_validated_excel(
    file: UploadFile,
    *,
    allowed_extensions: tuple[str, ...] = (".xlsx", ".xls"),
) -> bytes:
    filename = Path(file.filename or "").name
    extension = Path(filename).suffix.lower()
    if extension not in allowed_extensions:
        accepted = ", ".join(allowed_extensions)
        raise HTTPException(status_code=400, detail=f"Only {accepted} files are accepted.")
    file.filename = filename

    # Documented Vercel Production Limit: Vercel serverless functions have a strict 4.5 MB request body limit.
    # We validate this upfront to avoid opaque 500 crashes and return a clean 413 Payload Too Large.
    contents = await file.read(settings.MAX_UPLOAD_BYTES + 1)
    if len(contents) > settings.MAX_UPLOAD_BYTES:
        limit_mb = settings.MAX_UPLOAD_BYTES / (1024 * 1024)
        raise HTTPException(status_code=413, detail=f"Upload exceeds the {limit_mb:g} MB limit. Vercel production limits request bodies to 4.5 MB.")
    if not contents:
        raise HTTPException(status_code=400, detail="The uploaded file is empty.")

    valid_signature = (
        extension == ".xlsx" and contents.startswith(_XLSX_SIGNATURES)
    ) or (
        extension == ".xls" and contents.startswith(_XLS_SIGNATURE)
    )
    if not valid_signature:
        raise HTTPException(status_code=400, detail="The uploaded file is not a valid Excel workbook.")
    if extension == ".xlsx":
        _validate_xlsx_structure(contents)
    return contents
